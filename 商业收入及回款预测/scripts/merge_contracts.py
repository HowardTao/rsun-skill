#!/usr/bin/env python3
"""执行+未执行合同清单合并器（续签标记，收入/回款预测基础表）。

用法:
    python merge_contracts.py --active /tmp/active_contracts_<基准日>.json \
        --future /tmp/future_contracts_<基准日>.json \
        --base-date 2026-09-07 --store 弘阳家居南京江北店 \
        --output <输出目录>/执行+未执行合同清单_<基准日>.xlsx [--snapshot-time "..."]

输入为两个 guancli ds preview -f json 输出（与 make_report.py 同源）：
- --active: 在执行合同（起租日 ≤ 基准日 ≤ 到期日）
- --future: 未执行合同（起租日 > 基准日；本脚本自动应用 起租日 < 到期日 过滤，
  剔除起租日≥到期日或日期缺失/无效的记录，与 make_report.py --mode future 口径一致）

续签判定：同一（门店名称, 铺位号）下，在执行合同与未执行合同同时存在即视为续签：
- 在执行合同（原合同）: 是否续签=是, 续签描述="合同已续签，续签合同号：<未执行合同号>"
- 未执行合同（新合同）: 是否续签=否, 续签描述="该合同为续签合同，续签自合同号：<在执行合同号>"
- 其余合同: 是否续签=否, 续签描述为空
- 同铺位多份合同时按起租日升序列出全部合同号（顿号分隔）；铺位号为空不参与匹配。

输出为单 sheet 纯明细临时表：28 列 = 26 列基础列（含租金/物业费/营销推广费
总应收及优惠后总应收 + 各自的日单价/日净单价计算列，元/㎡/天）+ 是否续签/续签描述，
按铺位号、起租日、合同号排序（同铺位下在执行合同起租日必然早于未执行合同，
续签合同对自然相邻），作为后续收入和回款预测的基础表（临时文件，最终输出表另行提供）。
"""
import argparse
import json
import os
import sys
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_report import OUTPUT_COLS, calc_unit_prices, render_detail_sheet  # noqa: E402

MERGED_COLS = OUTPUT_COLS + ["是否续签", "续签描述"]
SHEET_NAME = "执行+未执行合同清单"
EXTRA_WIDTHS = {"是否续签": 10, "续签描述": 48}


def load_json(path, label):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, list):
        sys.exit(f"{label}输入 JSON 不是行记录数组，请确认是 guancli ds preview -f json 的输出: {path}")
    return data


def pos_key(row):
    """匹配键：(门店名称, 铺位号)；铺位号为空返回 None（不参与匹配）。"""
    store = str(row.get("门店名称") or "").strip()
    pos = str(row.get("铺位号") or "").strip()
    return (store, pos) if pos else None


def ids_sorted(rows):
    """取合同号列表，按起租日升序；跳过无合同号的行。"""
    return [str(r.get("合同号") or "").strip()
            for r in sorted(rows, key=lambda x: str(x.get("起租日") or ""))
            if str(r.get("合同号") or "").strip()]


def sort_key(row):
    return (
        str(row.get("铺位号") or ""),
        str(row.get("起租日") or "")[:10],
        str(row.get("合同号") or ""),
    )


def build(active_path, future_path, base_date_str, store, output, snapshot_time):
    active = load_json(active_path, "在执行合同")
    future = load_json(future_path, "未执行合同")
    try:
        date.fromisoformat(base_date_str)
    except ValueError:
        sys.exit(f"基准日无效: {base_date_str}，需要 YYYY-MM-DD 格式的真实日历日期")

    # 模式 B 附加口径：起租日 < 到期日（与 make_report.py --mode future 一致，本地过滤）
    kept, dropped = [], []
    for r in future:
        s, e = str(r.get("起租日") or "")[:10], str(r.get("到期日") or "")[:10]
        try:
            ok = bool(s) and bool(e) and date.fromisoformat(s) < date.fromisoformat(e)
        except ValueError:
            ok = False
        (kept if ok else dropped).append(r)
    future = kept
    calc_unit_prices(active)
    calc_unit_prices(future)
    if dropped:
        print(f"已过滤: 剔除{len(dropped)}行不满足 起租日<到期日（含起租日/到期日缺失或无效）的记录")
        for r in dropped[:10]:
            print(f"  - 合同号 {r.get('合同号')}, 起租日 {str(r.get('起租日') or '')[:10]}, "
                  f"到期日 {str(r.get('到期日') or '')[:10]}")

    act_by_pos = defaultdict(list)
    for r in active:
        k = pos_key(r)
        if k:
            act_by_pos[k].append(r)
    fut_by_pos = defaultdict(list)
    for r in future:
        k = pos_key(r)
        if k:
            fut_by_pos[k].append(r)

    rows = []
    renewed_act = renewed_fut = 0
    for r in active:
        row = dict(r)
        fut_ids = ids_sorted(fut_by_pos.get(pos_key(r), [])) if pos_key(r) else []
        if fut_ids:
            renewed_act += 1
            row["是否续签"] = "是"
            row["续签描述"] = f"合同已续签，续签合同号：{'、'.join(fut_ids)}"
        else:
            row["是否续签"] = "否"
            row["续签描述"] = ""
        rows.append(row)
    for r in future:
        row = dict(r)
        act_ids = ids_sorted(act_by_pos.get(pos_key(r), [])) if pos_key(r) else []
        if act_ids:
            renewed_fut += 1
            row["是否续签"] = "否"
            row["续签描述"] = f"该合同为续签合同，续签自合同号：{'、'.join(act_ids)}"
        else:
            row["是否续签"] = "否"
            row["续签描述"] = ""
        rows.append(row)

    rows.sort(key=sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    render_detail_sheet(ws, rows, MERGED_COLS)
    for j, col in enumerate(MERGED_COLS, 1):
        if col in EXTRA_WIDTHS:
            ws.column_dimensions[get_column_letter(j)].width = EXTRA_WIDTHS[col]
    wb.save(output)

    print(f"已保存: {output}")
    print(f"合并: 在执行{len(active)}行 + 未执行{len(future)}行 = {len(rows)}行, "
          f"门店: {store}, 基准日: {base_date_str}"
          + (f", 快照更新: {snapshot_time}" if snapshot_time else ""))
    print(f"续签: 在执行侧已续签{renewed_act}份, 未执行侧为续签合同{renewed_fut}份, "
          f"非续签未执行合同{len(future) - renewed_fut}份")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--active", required=True, help="在执行合同 JSON（guancli ds preview -f json）")
    ap.add_argument("--future", required=True, help="未执行合同 JSON（guancli ds preview -f json）")
    ap.add_argument("--base-date", required=True, help="基准日 YYYY-MM-DD")
    ap.add_argument("--store", required=True, help="门店全称（用于摘要打印）")
    ap.add_argument("--output", required=True, help="输出 xlsx 路径")
    ap.add_argument("--snapshot-time", default="", help="数据集快照更新时间（可选）")
    args = ap.parse_args()
    build(args.active, args.future, args.base_date, args.store, args.output,
          args.snapshot_time)


if __name__ == "__main__":
    main()
