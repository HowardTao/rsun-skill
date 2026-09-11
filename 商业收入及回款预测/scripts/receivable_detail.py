#!/usr/bin/env python3
"""合同全周期应收明细临时表生成 + 合并清单收款周期更新器（中间临时表，非最终输出表）。

用法:
    python receivable_detail.py \
        --active /tmp/active_contracts_<基准日>.json \
        --future /tmp/future_contracts_<基准日>.json \
        --receivable "/tmp/receivable_batches_<基准日>/*.json" \
        --base-date <基准日YYYY-MM-DD> --store <门店全称> \
        --output-detail <输出目录>/<门店简称>合同全周期应收明细_<基准日>.xlsx \
        --output-merged <输出目录>/执行+未执行合同清单_<基准日>.xlsx \
        [--snapshot-time "..."]
输入:
- --active / --future: 在执行/未执行合同清单 JSON（guancli ds preview -f json，
  与 merge_contracts.py 同源同口径；future 的 起租日<到期日 过滤在本脚本内部应用）。
- --receivable: 应收明细 JSON 的 glob（按合同号 IN 分批取自
  ADS-租费分析-合同全周期应收明细 ud1f309f6e04d4285b7364c4），
  一行一条账单费用区间明细；费用类型四类（租金/物业费/营销推广费/其他费用）
  全部纳入，按合同号与清单关联，仅保留清单内合同的行。

输出 1（--output-detail）应收明细临时表，单 sheet 纯明细 31 列:
- 合同标识 + 合同属性 16 列: 门店名称/合同号/铺位号/租户名称 + 区域/楼栋/品牌/业态/
  主品类/楼层/计租面积/起租日/到期日/铺位类型/合同状态/是否续签
  （12 个合同属性列均按合同号取自合并清单；区域/楼栋来自台账字段、随清单输入）
- 账单标识 2 列: 应收编号/账单主键
- 费项 3 列: 费用名/费用类型/税率
- 时间周期 3 列: 费用开始日/费用截止日/应收日
- 含税金额 5 列: 应收金额/优惠金额/优惠后应收金额/已收金额/未收金额
- 不含税计算 2 列: 不含税应收金额 = 应收金额/(1+税率)、
  不含税优惠后应收金额 = 优惠后应收金额/(1+税率)；税率缺失时留空；保留 2 位小数。
按 合同号/费用类型/费用开始日/费用名/应收编号 排序。

输出 2（--output-merged）合并清单升级版，单 sheet 纯明细 34 列 =
merge_contracts.py 的 30 列（含区域、楼栋、续签标记）+ 最后账期费用开始日
+ 最后账期费用截止日 + 支付周期 + 提前收款天数:
- 最后账期口径（2026-09-07 用户确认）: 费用类型=租金 的账单中
  费用截止日最晚的一期（截止日相同时取费用开始日更晚，再同取应收日更晚）。
- 最后账期费用开始日/最后账期费用截止日 = 该期账单的费用开始日/费用截止日。
- 支付周期 = 费用截止日 - 费用开始日 + 1（天，整数）。
- 提前收款天数 = 费用开始日 - 应收日（天，整数，负数表示后付，如实展示）。
- 无租金账单的合同四列留空；日期缺失/无效时对应列留空。
"""
import argparse
import glob as globlib
import os
import sys
from collections import Counter, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from make_report import check_store, num  # noqa: E402
from merge_contracts import (  # noqa: E402
    EXTRA_WIDTHS,
    MERGED_COLS,
    build_merged_rows,
    load_json,
)

DETAIL_SHEET = "合同全周期应收明细"
# 合同属性列（2026-09-08 新增 12 列：区域/楼栋取自台账映射，其余取自合并清单，按合同号关联）
ATTR_COLS = ["区域", "楼栋", "品牌", "业态", "主品类", "楼层", "计租面积",
             "起租日", "到期日", "铺位类型", "合同状态", "是否续签"]
DETAIL_COLS = [
    "门店名称", "合同号", "铺位号", "租户名称",
] + ATTR_COLS + [
    "应收编号", "账单主键",
    "费用名", "费用类型", "税率", "费用开始日", "费用截止日", "应收日",
    "应收金额", "优惠金额", "优惠后应收金额", "已收金额", "未收金额",
    "不含税应收金额", "不含税优惠后应收金额",
]
# 合并清单收款周期升级版列（原 MERGED32_COLS，2026-09-08 新增区域/楼栋列后为 34 列）
MERGED_PAY_COLS = MERGED_COLS + [
    "最后账期费用开始日", "最后账期费用截止日", "支付周期", "提前收款天数",
]

NUM_COLS = {
    "税率", "应收金额", "优惠金额", "优惠后应收金额", "已收金额", "未收金额",
    "不含税应收金额", "不含税优惠后应收金额",
    "租金总应收", "优惠后租金总应收", "物业费总应收", "优惠后物业费总应收",
    "营销推广费总应收", "优惠后营销推广费总应收",
    "租金日单价", "租金日净单价", "物业费日单价", "物业费日净单价",
    "营销推广费日单价", "营销推广费日净单价", "计租面积",
}
INT_COLS = {"支付周期", "提前收款天数"}
DATE_COLS = {"起租日", "到期日", "费用开始日", "费用截止日", "应收日",
             "最后账期费用开始日", "最后账期费用截止日"}
TEXT_COLS = {"合同号", "应收编号", "账单主键", "铺位号", "是否续签"}
WIDTHS = {
    "门店名称": 18, "合同号": 22, "铺位号": 13, "租户名称": 26, "应收编号": 28,
    "账单主键": 20, "费用名": 16, "费用类型": 10, "税率": 8, "费用开始日": 12,
    "费用截止日": 12, "应收日": 12, "应收金额": 13, "优惠金额": 12,
    "优惠后应收金额": 14, "已收金额": 13, "未收金额": 12, "不含税应收金额": 14,
    "不含税优惠后应收金额": 16, "支付周期": 10, "提前收款天数": 12,
    "最后账期费用开始日": 14, "最后账期费用截止日": 14,
    "区域": 12, "楼栋": 10, "品牌": 16, "业态": 10, "主品类": 12, "楼层": 8,
    "计租面积": 11, "起租日": 12, "到期日": 12, "铺位类型": 10, "合同状态": 10,
}

H_FONT = Font(name="PingFang SC", size=11, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor="1F4E79")
B_FONT = Font(name="PingFang SC", size=11)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def parse_date(v):
    s = str(v or "")[:10]
    if not s:
        return None
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


def render_sheet(ws, data, cols):
    """渲染明细 sheet：表头 + 数据行；冻结首行、自动筛选、列宽与数字格式。"""
    ws.freeze_panes = "A2"
    for j, col in enumerate(cols, 1):
        cell = ws.cell(1, j, col)
        cell.font = H_FONT
        cell.fill = H_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(data):
        for j, col in enumerate(cols, 1):
            v = row.get(col)
            if col in DATE_COLS and v:
                v = str(v)[:10]
            elif col in NUM_COLS:
                v = num(v)
            elif col in INT_COLS:
                v = num(v)
                v = int(v) if v is not None else None
            cell = ws.cell(i + 2, j, v)
            cell.font = B_FONT
            cell.border = BORDER
            if col in NUM_COLS and v is not None:
                cell.number_format = "#,##0.00"
            elif col in INT_COLS and v is not None:
                cell.number_format = "#,##0"
            elif col in TEXT_COLS:
                cell.number_format = "@"
    for j, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(
            col, EXTRA_WIDTHS.get(col, 12))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data) + 1}"


def load_receivable(pattern, valid_ids):
    """加载各批应收明细 JSON：合并、去重（账单主键优先）、过滤到清单合同号。"""
    paths = sorted(globlib.glob(pattern))
    if not paths:
        sys.exit(f"未匹配到应收明细 JSON: {pattern}")
    rows, seen = [], {}
    for p in paths:
        for r in load_json(p, "应收明细"):
            cid = str(r.get("合同号") or "").strip()
            if cid not in valid_ids:
                continue
            key = str(r.get("账单主键") or "").strip() or (
                f"{r.get('应收编号')}|{r.get('费用名')}|"
                f"{str(r.get('费用开始日') or '')[:10]}|{str(r.get('费用截止日') or '')[:10]}")
            if key in seen:
                continue
            seen[key] = True
            rows.append(r)
    return rows, paths


def calc_ex_tax(rows):
    """就地计算不含税金额两列：不含税 = 含税/(1+税率)，税率缺失时留空。"""
    for r in rows:
        t = num(r.get("税率"))
        for src, dst in (("应收金额", "不含税应收金额"),
                         ("优惠后应收金额", "不含税优惠后应收金额")):
            a = num(r.get(src))
            r[dst] = round(a / (1 + t), 2) if (a is not None and t is not None) else None


def calc_pay_period(rows):
    """每合同最后账期（费用类型=租金、费用截止日最晚，tie 取费用开始日/应收日更晚者）。

    返回 {合同号: (最后账期费用开始日, 最后账期费用截止日, 支付周期天数, 提前收款天数)}；
    日期为 YYYY-MM-DD 字符串；日期缺失/无效时对应值为 None。
    """
    best = {}
    for r in rows:
        if str(r.get("费用类型") or "").strip() != "租金":
            continue
        cid = str(r.get("合同号") or "").strip()
        if not cid:
            continue
        key = (parse_date(r.get("费用截止日")) or date.min,
               parse_date(r.get("费用开始日")) or date.min,
               parse_date(r.get("应收日")) or date.min)
        if cid not in best or key > best[cid][0]:
            best[cid] = (key, r)
    out = {}
    for cid, (key, r) in best.items():
        end, start, recv = key
        period = ((end - start).days + 1) if (end > date.min and start > date.min) else None
        advance = (start - recv).days if (start > date.min and recv > date.min) else None
        s_str = start.isoformat() if start > date.min else None
        e_str = end.isoformat() if end > date.min else None
        out[cid] = (s_str, e_str, period, advance)
    return out


def build(active_path, future_path, receivable_glob, base_date_str, store,
          output_detail, output_merged, snapshot_time):
    try:
        date.fromisoformat(base_date_str)
    except ValueError:
        sys.exit(f"基准日无效: {base_date_str}，需要 YYYY-MM-DD 格式的真实日历日期")

    active = load_json(active_path, "在执行合同")
    future = load_json(future_path, "未执行合同")
    merged_rows, renewed_act, renewed_fut, dropped = build_merged_rows(active, future)
    valid_ids = {str(r.get("合同号") or "").strip() for r in merged_rows
                 if str(r.get("合同号") or "").strip()}
    if dropped:
        print(f"已过滤(未执行清单): 剔除{len(dropped)}行不满足 起租日<到期日 的记录")

    check_store(merged_rows, store, "合并清单")

    receivable, paths = load_receivable(receivable_glob, valid_ids)
    calc_ex_tax(receivable)

    # 按合同号把 12 个合同属性列（区域/楼栋/品牌/.../是否续签）填充到应收明细行
    attr_by_cid = {}
    for row in merged_rows:
        cid = str(row.get("合同号") or "").strip()
        if cid:
            attr_by_cid[cid] = {col: row.get(col) for col in ATTR_COLS}
    for r in receivable:
        cid = str(r.get("合同号") or "").strip()
        r.update(attr_by_cid.get(cid, {}))

    receivable.sort(key=lambda r: (
        str(r.get("合同号") or ""), str(r.get("费用类型") or ""),
        str(r.get("费用开始日") or "")[:10], str(r.get("费用名") or ""),
        str(r.get("应收编号") or "")))

    covered = {str(r.get("合同号") or "").strip() for r in receivable}
    no_detail = sorted(valid_ids - covered)

    # --- 输出 1：应收明细临时表 ---
    wb = Workbook()
    ws = wb.active
    ws.title = DETAIL_SHEET
    render_sheet(ws, receivable, DETAIL_COLS)
    os.makedirs(os.path.dirname(os.path.abspath(output_detail)), exist_ok=True)
    wb.save(output_detail)

    # --- 输出 2：合并清单 34 列（+最后账期起止日/支付周期/提前收款天数）---
    pay = calc_pay_period(receivable)
    no_rent = sorted(cid for cid in valid_ids if cid not in pay)
    for row in merged_rows:
        cid = str(row.get("合同号") or "").strip()
        last_s, last_e, period, advance = pay.get(cid, (None, None, None, None))
        row["最后账期费用开始日"] = last_s
        row["最后账期费用截止日"] = last_e
        row["支付周期"] = period
        row["提前收款天数"] = advance
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "执行+未执行合同清单"
    render_sheet(ws2, merged_rows, MERGED_PAY_COLS)
    os.makedirs(os.path.dirname(os.path.abspath(output_merged)), exist_ok=True)
    wb2.save(output_merged)

    # --- 校验与摘要 ---
    print(f"已保存: {output_detail}")
    print(f"已保存: {output_merged}")
    print(f"门店: {store}, 基准日: {base_date_str}"
          + (f", 快照更新: {snapshot_time}" if snapshot_time else ""))
    print(f"应收明细输入: {len(paths)}个分批JSON, 合并去重后{len(receivable)}行, "
          f"覆盖清单合同{len(covered)}/{len(valid_ids)}份")
    if no_detail:
        print(f"清单中无应收明细的合同{len(no_detail)}份: {', '.join(no_detail[:10])}"
              + ("..." if len(no_detail) > 10 else ""))
    print(f"合并清单: {len(merged_rows)}行(34列) = 在执行{len(active)} + "
          f"未执行{len(future) - len(dropped)}, 续签: 在执行侧{renewed_act}份 / "
          f"未执行侧{renewed_fut}份")
    if no_rent:
        print(f"无租金账单（最后账期四列留空）的合同{len(no_rent)}份"
              + (f": {', '.join(no_rent[:10])}" + ("..." if len(no_rent) > 10 else "")))

    ft = Counter(r.get("费用类型") or "未分类" for r in receivable)
    print("应收明细费用类型分布:", dict(ft))
    sums = {c: sum(num(r.get(c)) or 0 for r in receivable) for c in
            ("应收金额", "优惠金额", "优惠后应收金额", "不含税应收金额",
             "不含税优惠后应收金额", "已收金额", "未收金额")}
    print("金额合计(元): " + ", ".join(f"{k}={v:,.2f}" for k, v in sums.items()))

    periods = sorted(v[2] for v in pay.values() if v[2] is not None)
    advs = sorted(v[3] for v in pay.values() if v[3] is not None)
    if periods:
        pc = Counter(periods)
        print(f"支付周期(天): 合同{len(periods)}份, 常见值: "
              + ", ".join(f"{d}天×{n}份" for d, n in pc.most_common(8)))
    if advs:
        neg = sum(1 for a in advs if a < 0)
        mid = advs[len(advs) // 2]
        print(f"提前收款天数: 合同{len(advs)}份, 最小{advs[0]} / 中位数{mid} / "
              f"最大{advs[-1]}天, 其中后付(负值){neg}份")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--active", required=True, help="在执行合同 JSON（guancli ds preview -f json）")
    ap.add_argument("--future", required=True, help="未执行合同 JSON（guancli ds preview -f json）")
    ap.add_argument("--receivable", required=True,
                    help="应收明细分批 JSON 的 glob（如 '/tmp/receivable_batches_日期/*.json'）")
    ap.add_argument("--base-date", required=True, help="基准日 YYYY-MM-DD")
    ap.add_argument("--store", required=True, help="门店全称（用于摘要打印）")
    ap.add_argument("--output-detail", required=True, help="应收明细临时表 xlsx 输出路径")
    ap.add_argument("--output-merged", required=True, help="合并清单（34 列）xlsx 输出路径")
    ap.add_argument("--snapshot-time", default="", help="数据集快照更新时间（可选）")
    args = ap.parse_args()
    build(args.active, args.future, args.receivable, args.base_date, args.store,
          args.output_detail, args.output_merged, args.snapshot_time)


if __name__ == "__main__":
    main()
