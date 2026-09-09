#!/usr/bin/env python3
"""合同临时表生成器（在执行/未执行合同，均为中间临时表，非最终输出表）。

用法:
    # 在执行合同临时表（起租日 ≤ 基准日 ≤ 到期日）
    python make_report.py --mode active --input <guancli导出的JSON> --base-date <基准日YYYY-MM-DD> \
        --store <门店全称> --output <输出.xlsx> [--snapshot-time "..."]

    # 未执行合同临时表（起租日 > 基准日，收入/回款预测输入）
    python make_report.py --mode future --input <guancli导出的JSON> --base-date <基准日YYYY-MM-DD> \
        --store <门店全称> --output <输出.xlsx> [--snapshot-time "..."]

输入 JSON 为 `guancli ds preview k3a8a2772ee4143d694a2ecd --filter ... -f json` 的输出
（行记录数组，一行一合同）。
两模式输出规格一致：单 sheet 纯明细临时表，输出列统一 28 列（OUTPUT_COLS）：
16 基础列（含区域、楼栋）+ 6 费用列（数据集现成字段）+ 6 计算列（日单价，元/㎡/天，
日单价 = 总应收/(到期日-起租日+1)/计租面积；日净单价 = 优惠后总应收/同口径；
计租面积或日期缺失/无效、天数为非正数时留空），
无汇总页；作为后续收入/回款预测的中间基表，最终输出表另行提供。
口径:
- 在执行(active) = 起租日 <= 基准日 <= 到期日(日期窗口判定),含正铺+多经;合同状态仅展示。
- 未执行(future) = 起租日 > 基准日(已签约尚未起租);合同状态仅展示。
"""
import argparse
import json
import os
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_COLS = [
    "门店名称", "区域", "楼栋", "合同号", "铺位号", "租户名称", "品牌", "业态",
    "主品类", "楼层", "计租方式", "计租面积", "起租日", "到期日", "铺位类型", "合同状态",
    "租金总应收", "优惠后租金总应收", "物业费总应收", "优惠后物业费总应收",
    "营销推广费总应收", "优惠后营销推广费总应收",
    "租金日单价", "租金日净单价", "物业费日单价", "物业费日净单价",
    "营销推广费日单价", "营销推广费日净单价",
]
# 计算列 = (总应收列, 优惠后总应收列, 日单价列, 日净单价列)
PRICE_PAIRS = [
    ("租金总应收", "优惠后租金总应收", "租金日单价", "租金日净单价"),
    ("物业费总应收", "优惠后物业费总应收", "物业费日单价", "物业费日净单价"),
    ("营销推广费总应收", "优惠后营销推广费总应收", "营销推广费日单价", "营销推广费日净单价"),
]
NUM_COLS = {
    "计租面积", "租金总应收", "优惠后租金总应收", "物业费总应收",
    "优惠后物业费总应收", "营销推广费总应收", "优惠后营销推广费总应收",
    "租金日单价", "租金日净单价", "物业费日单价", "物业费日净单价",
    "营销推广费日单价", "营销推广费日净单价",
}
DATE_COLS = {"起租日", "到期日"}
WIDTHS = {
    "门店名称": 18, "区域": 12, "楼栋": 10, "合同号": 15, "铺位号": 13, "租户名称": 26,
    "品牌": 16, "业态": 10, "主品类": 12, "楼层": 8, "计租方式": 10, "计租面积": 11,
    "起租日": 12, "到期日": 12, "铺位类型": 10, "合同状态": 10,
    "租金总应收": 13, "优惠后租金总应收": 14, "物业费总应收": 13,
    "优惠后物业费总应收": 14, "营销推广费总应收": 14, "优惠后营销推广费总应收": 15,
    "租金日单价": 11, "租金日净单价": 12, "物业费日单价": 11, "物业费日净单价": 12,
    "营销推广费日单价": 13, "营销推广费日净单价": 14,
}
MODE_META = {
    "active": {"sheet": "在执行合同", "scope": "起租日≤基准日≤到期日（日期窗口判定）"},
    "future": {"sheet": "未执行合同", "scope": "起租日>基准日（已签约尚未起租）"},
}

H_FONT = Font(name="PingFang SC", size=11, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor="1F4E79")
B_FONT = Font(name="PingFang SC", size=11)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def clean(v, col):
    if col in DATE_COLS and v:
        return str(v)[:10]
    if col in NUM_COLS:
        return num(v)
    return v


def calc_unit_prices(rows):
    """就地计算 6 个日单价列（元/㎡/天）：
    日单价 = 总应收/(到期日-起租日+1)/计租面积；日净单价 = 优惠后总应收/同口径。
    计租面积缺失/为0、日期缺失/无效、天数为非正数时留空（None）。
    """
    for r in rows:
        s, e = str(r.get("起租日") or "")[:10], str(r.get("到期日") or "")[:10]
        area = num(r.get("计租面积"))
        days = None
        try:
            if s and e:
                days = (date.fromisoformat(e) - date.fromisoformat(s)).days + 1
        except ValueError:
            days = None
        denom = area * days if (area and days and days > 0) else None
        for gross, net, out_g, out_n in PRICE_PAIRS:
            g, n = num(r.get(gross)), num(r.get(net))
            r[out_g] = round(g / denom, 6) if (g is not None and denom) else None
            r[out_n] = round(n / denom, 6) if (n is not None and denom) else None


def render_detail_sheet(ws, data, cols):
    """渲染明细 sheet：第 1 行表头、第 2 行起数据；冻结首行并加自动筛选。"""
    ws.freeze_panes = "A2"
    for j, col in enumerate(cols, 1):
        cell = ws.cell(1, j, col)
        cell.font = H_FONT; cell.fill = H_FILL; cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(data):
        for j, col in enumerate(cols, 1):
            v = clean(row.get(col), col)
            cell = ws.cell(i + 2, j, v)
            cell.font = B_FONT; cell.border = BORDER
            if col in NUM_COLS and v is not None:
                cell.number_format = "#,##0.00"
            elif col == "合同号":
                cell.number_format = "@"
    for j, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(col, 12)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data) + 1}"


def build(input_path, base_date_str, store, output, snapshot_time, pos_note, mode):
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, list):
        sys.exit("输入 JSON 不是行记录数组，请确认是 guancli ds preview -f json 的输出")

    meta = MODE_META[mode]
    calc_unit_prices(data)

    # 模式 B 附加口径：只保留 起租日 < 到期日 的记录（guancli filter 不支持字段间比较，本地过滤）
    dropped = []
    if mode == "future":
        kept = []
        for r in data:
            s, e = str(r.get("起租日") or "")[:10], str(r.get("到期日") or "")[:10]
            try:
                ok = bool(s) and bool(e) and date.fromisoformat(s) < date.fromisoformat(e)
            except ValueError:
                ok = False
            (kept if ok else dropped).append(r)
        data = kept

    n = len(data)
    contracts = {str(r.get("合同号") or "").strip() for r in data if r.get("合同号")}
    area = sum(num(r.get("计租面积")) or 0 for r in data)
    starts = sorted(str(r.get("起租日") or "")[:10] for r in data if r.get("起租日"))
    ends = sorted(str(r.get("到期日") or "")[:10] for r in data if r.get("到期日"))
    pt = Counter(r.get("铺位类型") or "未分类" for r in data)
    st = Counter(r.get("合同状态") or "未分类" for r in data)

    wb = Workbook()
    ws = wb.active
    ws.title = meta["sheet"]
    render_detail_sheet(ws, data, OUTPUT_COLS)
    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    wb.save(output)

    print(f"已保存: {output}")
    print(f"模式: {mode}（{meta['scope']}）, 门店: {store}, 基准日: {base_date_str}, "
          f"铺位类型: {pos_note}"
          + (f", 快照更新: {snapshot_time}" if snapshot_time else ""))
    print(f"临时表: {n}行(唯一合同号{len(contracts)}), 计租面积{area:,.2f}㎡")
    if dropped:
        print(f"已过滤: 剔除{len(dropped)}行不满足 起租日<到期日（含起租日/到期日缺失或无效）的记录")
        for r in dropped[:10]:
            print(f"  - 合同号 {r.get('合同号')}, 起租日 {str(r.get('起租日') or '')[:10]}, "
                  f"到期日 {str(r.get('到期日') or '')[:10]}")
    if starts:
        print(f"起租日范围: {starts[0]} ~ {starts[-1]}")
    if ends:
        print(f"到期日范围: {ends[0]} ~ {ends[-1]}")
    print("铺位类型分布:", dict(pt))
    print("合同状态分布:", dict(st))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="guancli ds preview -f json 输出文件")
    ap.add_argument("--base-date", required=True, help="基准日 YYYY-MM-DD")
    ap.add_argument("--store", required=True, help="门店全称（用于摘要打印）")
    ap.add_argument("--output", required=True, help="输出 xlsx 路径")
    ap.add_argument("--snapshot-time", default="", help="数据集快照更新时间（可选）")
    ap.add_argument("--pos-note", default="正铺+多经", help="铺位类型范围说明")
    ap.add_argument("--mode", choices=["active", "future"], default="active",
                    help="active=在执行合同(起租日≤基准日≤到期日); "
                         "future=未执行合同(起租日>基准日)")
    args = ap.parse_args()
    try:
        date.fromisoformat(args.base_date)
    except ValueError:
        sys.exit(f"基准日无效: {args.base_date}，需要 YYYY-MM-DD 格式的真实日历日期")
    build(args.input, args.base_date, args.store, args.output, args.snapshot_time,
          args.pos_note, args.mode)


if __name__ == "__main__":
    main()
